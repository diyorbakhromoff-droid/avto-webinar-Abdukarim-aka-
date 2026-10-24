"""Dashboard, statistika, reklama xarajati, jonli webinar natijalari."""
import io
from datetime import date, timedelta

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.admin.chain import stats_for
from app.api.common import funnel_or_404, iso
from app.db import get_db
from app.models import AdSpend, ChainMessage, KeyValue, Lead, Payment, Question, ScheduledMessage
from app.security.auth import ANY, CONTENT, SALES
from app.services import stats as st
from app.timeutil import now_utc, to_local

router = APIRouter(prefix="/api/admin", tags=["admin"])


def flt_dep(funnel_id: int, date_from: date | None = None, date_to: date | None = None, source: str | None = None,
            slot: str | None = None, weekday: int | None = None, method: str | None = None,
            manager_id: int | None = None) -> st.StatsFilter:
    return st.StatsFilter(funnel_id, date_from, date_to, source or None, slot or None, weekday, method or None,
                          manager_id)


@router.get("/dashboard")
async def dashboard(funnel_id: int, period: str = "today", db: AsyncSession = Depends(get_db), _=Depends(ANY)):
    await funnel_or_404(db, funnel_id)
    today = to_local(now_utc()).date()
    start = {"today": today, "7d": today - timedelta(days=6), "30d": today - timedelta(days=29)}.get(period, today)
    n = await st.funnel_numbers(db, st.StatsFilter(funnel_id, start, today))
    cps = round(n["spend"] / n["full"]) if n["spend"] and n["full"] else None
    labels = [("Start", "start"), ("Telefon berdi", "phone"), ("Sessiya tanladi", "session"),
              ("1-kunni ko'rdi", "day1"), ("2-kun, ofertagacha", "offer"), ("Tugma bosdi", "clicked"),
              ("Chek yubordi", "receipt"), ("Boshlang'ich to'lov", "deposit"), ("To'liq to'lov", "full")]
    funnel, prev = [], None
    for label, key in labels:
        # to'lovlar 1-kunni ko'rganlarga nisbatan (TZ 10), qolganlari — oldingi bosqichga
        base = n["day1"] if key in ("deposit", "full") else prev
        funnel.append({"label": label, "n": n[key], "rate": st.rate(n[key], base) if base is not None else None})
        if key not in ("deposit", "full"):
            prev = n[key]
    pending = (await db.execute(select(func.count()).select_from(Payment).join(Lead, Lead.id == Payment.lead_id)
                                .where(Lead.funnel_id == funnel_id, Payment.status == "pending"))).scalar()
    unanswered = (await db.execute(select(func.count()).select_from(Question).join(Lead, Lead.id == Question.lead_id)
                                   .where(Lead.funnel_id == funnel_id, Question.answer.is_(None),
                                          Question.source == "bot"))).scalar()
    chat_new = (await db.execute(select(func.count()).select_from(Question).join(Lead, Lead.id == Question.lead_id)
                                 .where(Lead.funnel_id == funnel_id, Question.answer.is_(None),
                                        Question.source == "miniapp"))).scalar()
    failed = (await db.execute(select(func.count()).select_from(ScheduledMessage)
                               .join(Lead, Lead.id == ScheduledMessage.lead_id)
                               .where(Lead.funnel_id == funnel_id, ScheduledMessage.status == "failed",
                                      ScheduledMessage.reason.not_like("403%"),
                                      ScheduledMessage.send_at >= now_utc() - timedelta(days=7)))).scalar()
    return {
        "kpis": {"start": n["start"], "phone": n["phone"], "day1_watched": n["day1"], "full_paid": n["full"],
                 "deposit_paid": n["deposit"], "cost_per_sale": cps},
        "funnel": funnel,
        "upcoming": [{**u, "at": iso(u["at"])} for u in await st.upcoming(db, funnel_id)],
        "pending_receipts": pending, "unanswered_questions": unanswered, "unanswered_chat": chat_new,
        "failed_messages": failed,
    }


@router.get("/stats")
async def stats(flt: st.StatsFilter = Depends(flt_dep), db: AsyncSession = Depends(get_db), _=Depends(ANY)):
    f = await funnel_or_404(db, flt.funnel_id)
    n = await st.funnel_numbers(db, flt)
    msg_stats = await stats_for(db, f.id)
    msgs = (await db.execute(select(ChainMessage).where(ChainMessage.funnel_id == f.id)
                             .order_by(ChainMessage.sort, ChainMessage.id))).scalars()
    return {
        "metrics": st.metrics(n),
        "retention": await st.retention(db, flt, f),
        "messages": [{"id": m.id, "title": m.title, **msg_stats.get(m.id, {"sent": 0, "clicks": 0, "blocked_after": 0})}
                     for m in msgs],
        "steps": await st.steps(db, flt, f),
        "live_compare": await st.live_compare(db, f.id, n),
    }


@router.get("/funnels/{fid}/watch-stats")
async def watch_stats(fid: int, day: int = 1, step: int = 5, date_from: date | None = None,
                      date_to: date | None = None, source: str | None = None, slot: str | None = None,
                      db: AsyncSession = Depends(get_db), _=Depends(ANY)):
    """Videoning har 5 daqiqasida nechta odam ko'rib turibdi va qayerda ko'rish kamayadi."""
    f = await funnel_or_404(db, fid)
    flt = st.StatsFilter(fid, date_from, date_to, source or None, slot or None)
    return await st.watch_curve(db, flt, f, 2 if day == 2 else 1, max(1, min(step, 30)))


@router.get("/stats/export.xlsx")
async def stats_export(flt: st.StatsFilter = Depends(flt_dep), db: AsyncSession = Depends(get_db), _=Depends(ANY)):
    data = await stats(flt, db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ko'rsatkichlar"
    ws.append(["Ko'rsatkich", "Qiymat", "O'tish %"])
    for m in data["metrics"]:
        ws.append([m["label"], m.get("display") or m["value"], round(m["rate"] * 100, 1) if m["rate"] is not None else None])
    ws2 = wb.create_sheet("Xabarlar")
    ws2.append(["Xabar", "Yuborildi", "Tugma bosildi", "Keyin bloklagan"])
    for m in data["messages"]:
        ws2.append([m["title"], m["sent"], m["clicks"], m["blocked_after"]])
    ws3 = wb.create_sheet("Qadamlar")
    ws3.append(["Qadam", "To'xtadi", "Bot eslatmasidan qaytdi", "Qo'ng'iroqdan qaytdi"])
    for s in data["steps"]:
        ws3.append([s["label"], s["stuck"], s["returned_by_bot"], s["returned_by_call"]])
    for r in data["retention"]:
        w = wb.create_sheet(f"Tomosha {r['day']}-kun")
        w.append(["Daqiqa", "Qolgan ulush %"])
        for p in r["points"]:
            w.append([p["minute"], round(p["share"] * 100, 1)])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="statistika.xlsx"'})


class AdSpendIn(BaseModel):
    source: str
    date: date
    amount: int


@router.get("/funnels/{fid}/ad-spend")
async def ad_spend(fid: int, db: AsyncSession = Depends(get_db), _=Depends(ANY)):
    rows = (await db.execute(select(AdSpend).where(AdSpend.funnel_id == fid).order_by(AdSpend.date.desc()))).scalars()
    return [{"id": a.id, "source": a.source, "date": a.date.isoformat(), "amount": a.amount} for a in rows]


@router.post("/funnels/{fid}/ad-spend")
async def add_ad_spend(fid: int, body: AdSpendIn, db: AsyncSession = Depends(get_db), _=Depends(SALES)):
    a = AdSpend(funnel_id=fid, **body.model_dump())
    db.add(a)
    await db.commit()
    return {"id": a.id, "source": a.source, "date": a.date.isoformat(), "amount": a.amount}


class LiveResultsIn(BaseModel):
    day1_watch_rate: float | None = None
    paid_rate: float | None = None
    cost_per_sale: int | None = None


@router.get("/funnels/{fid}/live-results")
async def live_results(fid: int, db: AsyncSession = Depends(get_db), _=Depends(ANY)):
    kv = await db.get(KeyValue, f"live_results:{fid}")
    return kv.value if kv else LiveResultsIn().model_dump()


@router.put("/funnels/{fid}/live-results")
async def put_live_results(fid: int, body: LiveResultsIn, db: AsyncSession = Depends(get_db), _=Depends(CONTENT)):
    key = f"live_results:{fid}"
    kv = await db.get(KeyValue, key)
    if kv:
        kv.value = body.model_dump()
    else:
        db.add(KeyValue(key=key, value=body.model_dump()))
    await db.commit()
    return body.model_dump()
